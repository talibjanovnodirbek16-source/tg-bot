import telebot
import os
import tempfile
import sqlite3
from groq import Groq

TELEGRAM_TOKEN = "8646600794:AAFVUJUIZS5Mq8AHnoBibTMxx0iPRBbuksk"
GROQ_API_KEY = "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"

client = Groq(api_key=GROQ_API_KEY)
bot = telebot.TeleBot(TELEGRAM_TOKEN)

def init_db():
    conn = sqlite3.connect("/tmp/costs.db")
    c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS costs (
        article TEXT PRIMARY KEY,
        cost REAL
    )""")
    conn.commit()
    conn.close()

def get_costs():
    conn = sqlite3.connect("/tmp/costs.db")
    c = conn.cursor()
    c.execute("SELECT article, cost FROM costs")
    rows = c.fetchall()
    conn.close()
    return {row[0]: row[1] for row in rows}

def save_cost(article, cost):
    conn = sqlite3.connect("/tmp/costs.db")
    c = conn.cursor()
    c.execute("INSERT OR REPLACE INTO costs (article, cost) VALUES (?, ?)", (article, cost))
    conn.commit()
    conn.close()

init_db()
user_states = {}

def ask_ai(text):
    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": "Ты финансовый помощник для маркетплейса Wildberries. Отвечай на русском языке."},
            {"role": "user", "content": text}
        ]
    )
    return response.choices[0].message.content

def process_wb_report(chat_id, downloaded):
    import openpyxl
    costs = get_costs()

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        f.write(downloaded)
        tmp_path = f.name

    wb = openpyxl.load_workbook(tmp_path)
    os.unlink(tmp_path)

    if "Поартикульно" not in wb.sheetnames:
        bot.send_message(chat_id, "❌ Не нашёл лист 'Поартикульно'. Отправь правильный отчёт WB.")
        return

    ws = wb["Поартикульно"]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]

    try:
        art_idx = headers.index("Артикул поставщика")
        sales_idx = headers.index("Продажи, шт")
        revenue_idx = headers.index("К перечислению")
        logistics_idx = headers.index("Логистика")
    except:
        bot.send_message(chat_id, "❌ Не могу найти нужные колонки в файле.")
        return

    report = "📊 *ФИНАНСОВЫЙ ОТЧЁТ WB*\n\n"
    missing_costs = []
    total_profit = 0
    total_revenue = 0

    for row in rows[1:]:
        article = row[art_idx]
        if not article:
            continue
        sales = row[sales_idx] or 0
        revenue = row[revenue_idx] or 0
        logistics = row[logistics_idx] or 0

        cost = costs.get(str(article))
        if cost is None:
            missing_costs.append(str(article))
            continue

        total_cost = cost * sales
        profit = revenue - total_cost - logistics
        margin = (profit / revenue * 100) if revenue > 0 else 0
        roi = (profit / total_cost * 100) if total_cost > 0 else 0

        total_profit += profit
        total_revenue += revenue

        report += f"▪️ *{article}*\n"
        report += f"   Продажи: {int(sales)} шт\n"
        report += f"   К перечислению: {revenue:,.0f} ₽\n"
        report += f"   Себестоимость: {total_cost:,.0f} ₽\n"
        report += f"   Прибыль: {profit:,.0f} ₽\n"
        report += f"   Маржа: {margin:.1f}% | ROI: {roi:.1f}%\n\n"

    total_margin = (total_profit / total_revenue * 100) if total_revenue > 0 else 0
    report += f"━━━━━━━━━━━━━━━\n"
    report += f"💰 *ИТОГО*\n"
    report += f"   Выручка: {total_revenue:,.0f} ₽\n"
    report += f"   Прибыль: {total_profit:,.0f} ₽\n"
    report += f"   Маржа: {total_margin:.1f}%\n"

    if missing_costs:
        report += f"\n⚠️ Нет себестоимости для:\n"
        for a in missing_costs:
            report += f"   - {a}\n"
        report += "\nИспользуй /cost чтобы добавить."

    bot.send_message(chat_id, report, parse_mode="Markdown")

@bot.message_handler(commands=["start"])
def start(message):
    bot.reply_to(message, """Салам! Я финансовый помощник для WB.

📊 Отправь Excel отчёт — посчитаю прибыль, маржу и ROI
💰 /cost — добавить или обновить себестоимости
❓ Задай любой вопрос про WB""")

@bot.message_handler(commands=["cost"])
def cost_command(message):
    costs = get_costs()
    if not costs:
        bot.reply_to(message, """💰 Себестоимости не заданы.

Отправь в формате (каждый артикул с новой строки):
артикул: цена

Пример:
двойка_бежевый: 1350
двойка_белый: 1400""")
    else:
        text = "📋 *Текущие себестоимости:*\n\n"
        for art, cost in costs.items():
            text += f"▪️ {art}: {cost:,.0f} ₽\n"
        text += "\n✏️ Отправь новые данные чтобы обновить."
        bot.reply_to(message, text, parse_mode="Markdown")
    user_states[message.chat.id] = "entering_costs"

@bot.message_handler(content_types=["document"])
def handle_document(message):
    try:
        file_info = bot.get_file(message.document.file_id)
        downloaded = bot.download_file(file_info.file_path)
        file_name = message.document.file_name.lower()

        if file_name.endswith(".xlsx") or file_name.endswith(".xls"):
            bot.reply_to(message, "⏳ Читаю отчёт WB...")
            process_wb_report(message.chat.id, downloaded)

        elif file_name.endswith(".pdf"):
            bot.reply_to(message, "⏳ Читаю PDF...")
            import pdfplumber
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
                f.write(downloaded)
                tmp_path = f.name
            text = ""
            with pdfplumber.open(tmp_path) as pdf:
                for page in pdf.pages:
                    text += page.extract_text() or ""
            os.unlink(tmp_path)
            reply = ask_ai(f"Проанализируй документ:\n\n{text[:4000]}")
            bot.reply_to(message, reply)

        else:
            bot.reply_to(message, "Отправь Excel (.xlsx) отчёт из WB или PDF документ.")

    except Exception as e:
        bot.reply_to(message, f"Ошибка: {e}")

@bot.message_handler(func=lambda m: True)
def handle_text(message):
    chat_id = message.chat.id

    if user_states.get(chat_id) == "entering_costs":
        lines = message.text.strip().split("\n")
        saved = []
        errors = []
        for line in lines:
            if ":" in line:
                parts = line.split(":")
                article = parts[0].strip()
                try:
                    cost = float(parts[1].strip().replace(",", "."))
                    save_cost(article, cost)
                    saved.append(f"{article}: {cost:,.0f} ₽")
                except:
                    errors.append(line)
            else:
                errors.append(line)

        reply = ""
        if saved:
            reply += "✅ Сохранено:\n" + "\n".join(saved)
        if errors:
            reply += "\n❌ Ошибка в строках:\n" + "\n".join(errors)

        user_states[chat_id] = None
        bot.reply_to(message, reply or "Ничего не сохранено.")
        return

    try:
        reply = ask_ai(message.text)
        bot.reply_to(message, reply)
    except Exception as e:
        bot.reply_to(message, f"Ошибка: {e}")

bot.polling()
